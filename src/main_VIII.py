
import os
import time
import zipfile
import webbrowser
import unicodedata
from pathlib import Path
from io import BytesIO

import mysql.connector
from pypdf import PdfReader
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# PROGRAMA: MAIN VIII - RPA PROCESSADOR DE DOCUMENTOS
# ============================================================
# ETAPAS:
# 1. Conectar ao MySQL.
# 2. Abrir o Gmail para acesso manual.
# 3. Ler os PDFs do arquivo teste.zip.
# 4. Extrair e classificar o conteúdo.
# 5. Salvar os documentos no banco de dados.
# 6. Exportar os registros para Excel.
# 7. Exibir o resumo e encerrar a conexão.
#
# IMPORTANTE:
# O ZIP precisa estar na mesma pasta deste arquivo.
# A limpeza do banco NÃO é automática.
# ============================================================


# ------------------------------------------------------------
# CONFIGURAÇÕES
# ------------------------------------------------------------

PASTA_PROJETO = Path(__file__).resolve().parent
ARQUIVO_ZIP = PASTA_PROJETO / "teste.zip"
ARQUIVO_EXCEL = PASTA_PROJETO / "relatorio_emails.xlsx"

CONFIG_MYSQL = {
    "host": "localhost",
    "user": "root",
    "password": "",
    "database": "EmailBD",
    "port": 3306
}


# ------------------------------------------------------------
# CATEGORIAS E PALAVRAS-CHAVE
# ------------------------------------------------------------

CATEGORIAS_KEYWORDS = {
    "Trabalho": [
        "relatório", "indicadores", "absenteísmo",
        "atividades operacionais", "estratégica"
    ],
    "Estudos": [
        "arqueologia", "algoritmo", "paleontologia",
        "osteologia", "taekwondo", "primeiros socorros"
    ],
    "Tecnologia": [
        "programação", "python", "javascript", "sistema",
        "inteligência artificial", "tesla", "windows",
        "videogames", "copilotos", "robocop"
    ],
    "Entretenimento": [
        "filme", "transformers", "marvel", "ficção",
        "diaclone", "microman", "videogames"
    ]
}


# ------------------------------------------------------------
# NORMALIZAR TEXTO
# ------------------------------------------------------------

def normalizar_texto(texto):
    """Converte o texto para minúsculas e remove acentos."""

    texto = texto.lower()

    texto = unicodedata.normalize("NFD", texto)

    texto = "".join(
        caractere
        for caractere in texto
        if unicodedata.category(caractere) != "Mn"
    )

    return texto


# ------------------------------------------------------------
# CLASSIFICAR DOCUMENTO
# ------------------------------------------------------------

def classificar_email(texto):
    """Classifica o documento conforme as palavras-chave."""

    texto = normalizar_texto(texto)

    pontuacoes = {}

    for categoria, palavras in CATEGORIAS_KEYWORDS.items():
        pontuacoes[categoria] = 0

        for palavra in palavras:
            palavra = normalizar_texto(palavra)

            if palavra in texto:
                pontuacoes[categoria] += 1

    maior_pontuacao = max(pontuacoes.values(), default=0)

    if maior_pontuacao == 0:
        return "Não classificado"

    # Seleciona a categoria com maior pontuação.
    return max(pontuacoes, key=pontuacoes.get)


# ------------------------------------------------------------
# CONECTAR AO MYSQL
# ------------------------------------------------------------

def conectar_banco():
    """Estabelece a conexão com o banco EmailBD."""

    conexao = mysql.connector.connect(**CONFIG_MYSQL)

    if conexao.is_connected():
        return conexao

    raise ConnectionError("Não foi possível conectar ao MySQL.")


# ------------------------------------------------------------
# ABRIR O GMAIL
# ------------------------------------------------------------

def abrir_gmail():
    """Abre o Gmail e permite o login manual."""

    print("\nAbrindo o Gmail no navegador...")

    webbrowser.open("https://mail.google.com/", new=2)

    time.sleep(3)

    print("Faça o login manualmente, se necessário.")
    print("Não informe sua senha no terminal do Python.")

    input(
        "\nQuando terminar, pressione ENTER "
        "para continuar o processamento do ZIP..."
    )


# ------------------------------------------------------------
# PROCESSAR OS PDFs DO ZIP
# ------------------------------------------------------------

def processar_zip(conexao):
    """Lê, classifica e salva os PDFs dentro do ZIP."""

    processados = 0
    inseridos = 0
    ignorados = 0
    erros = 0

    if not ARQUIVO_ZIP.exists():
        raise FileNotFoundError(
            f"Arquivo ZIP não encontrado: {ARQUIVO_ZIP}"
        )

    cursor = conexao.cursor()

    try:
        with zipfile.ZipFile(ARQUIVO_ZIP, "r") as arquivo_zip:

            arquivos_pdf = [
                nome
                for nome in arquivo_zip.namelist()
                if nome.lower().endswith(".pdf")
                and not nome.endswith("/")
            ]

            print(f"\nPDFs encontrados no ZIP: {len(arquivos_pdf)}")

            for indice, nome_arquivo in enumerate(arquivos_pdf, start=1):

                print("\n" + "-" * 55)
                print(f"DOCUMENTO {indice}/{len(arquivos_pdf)}")
                print(Path(nome_arquivo).name)

                try:
                    # Lê o PDF diretamente do ZIP.
                    dados_pdf = arquivo_zip.read(nome_arquivo)

                    leitor = PdfReader(BytesIO(dados_pdf))

                    # Extrai o texto de todas as páginas.
                    partes_texto = []

                    for pagina in leitor.pages:
                        partes_texto.append(
                            pagina.extract_text() or ""
                        )

                    texto = "\n".join(partes_texto).strip()

                    if not texto:
                        print("PDF sem texto extraível. Ignorado.")
                        ignorados += 1
                        continue

                    processados += 1

                    # Usa o nome do PDF como assunto.
                    assunto = Path(nome_arquivo).stem

                    categoria = classificar_email(
                        assunto + "\n" + texto
                    )

                    # Verifica se o documento já foi cadastrado.
                    cursor.execute(
                        """
                        SELECT id
                        FROM tb_email
                        WHERE remetente = %s
                          AND assunto = %s
                        LIMIT 1
                        """,
                        ("ZIP local", assunto)
                    )

                    if cursor.fetchone():
                        print("Documento já cadastrado. Ignorado.")
                        ignorados += 1
                        continue

                    # Grava o documento no banco.
                    cursor.execute(
                        """
                        INSERT INTO tb_email (
                            remetente,
                            destinatario,
                            assunto,
                            conteudo,
                            data_hr,
                            filtragem,
                            status
                        )
                        VALUES (%s, %s, %s, %s, %s, %s, %s)
                        """,
                        (
                            "ZIP local",
                            "Não aplicável",
                            assunto,
                            texto,
                            None,
                            categoria,
                            "OK"
                        )
                    )

                    # Confirma a gravação.
                    conexao.commit()

                    inseridos += 1

                    print(f"Categoria: {categoria}")
                    print("Documento salvo no MySQL!")

                except Exception as erro_documento:
                    conexao.rollback()
                    erros += 1

                    print(
                        f"Erro ao processar este PDF: "
                        f"{erro_documento}"
                    )

    finally:
        cursor.close()

    return processados, inseridos, ignorados, erros


# ------------------------------------------------------------
# EXPORTAR DADOS PARA EXCEL
# ------------------------------------------------------------

def exportar_excel(conexao):
    """Exporta os registros do MySQL para uma planilha."""

    cursor = conexao.cursor()

    try:
        cursor.execute(
            """
            SELECT
                id,
                remetente,
                destinatario,
                assunto,
                conteudo,
                data_hr,
                filtragem,
                status
            FROM tb_email
            ORDER BY id
            """
        )

        registros = cursor.fetchall()

        colunas = [
            "ID",
            "Remetente",
            "Destinatário",
            "Assunto",
            "Conteúdo",
            "Data/Hora",
            "Categoria",
            "Status"
        ]

        # Cria uma nova planilha.
        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Relatorio de Emails"

        planilha.append(colunas)

        # Insere os dados consultados.
        for registro in registros:
            planilha.append(list(registro))

        # Formata o cabeçalho.
        for celula in planilha[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Ajusta a largura das colunas.
        larguras = [10, 24, 24, 42, 60, 22, 22, 14]

        for indice, largura in enumerate(larguras, start=1):
            planilha.column_dimensions[
                get_column_letter(indice)
            ].width = largura

        # Mantém o cabeçalho visível durante a rolagem.
        planilha.freeze_panes = "A2"

        # Ativa os filtros.
        planilha.auto_filter.ref = planilha.dimensions

        # Habilita a quebra de linha no conteúdo.
        for linha in planilha.iter_rows(min_row=2):
            for celula in linha:
                celula.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        workbook.save(ARQUIVO_EXCEL)

        print(f"\nExcel exportado: {ARQUIVO_EXCEL}")
        print(f"Registros exportados: {len(registros)}")

    finally:
        cursor.close()


# ------------------------------------------------------------
# FUNÇÃO PRINCIPAL
# ------------------------------------------------------------

def main():

    os.system("cls" if os.name == "nt" else "clear")

    print("=" * 60)
    print("       MAIN VIII - RPA PROCESSADOR DE DOCUMENTOS")
    print("=" * 60)

    conexao = None

    try:
        # 1. Conecta ao banco.
        print("\nConectando ao MySQL...")

        conexao = conectar_banco()

        print("Conexão realizada com sucesso!")

        # 2. Confirma que a tabela existe.
        cursor = conexao.cursor()

        try:
            cursor.execute("SELECT COUNT(*) FROM tb_email")
            total_atual = cursor.fetchone()[0]

        finally:
            cursor.close()

        print(f"Registros atuais no banco: {total_atual}")

        # 3. Abre o Gmail para acesso manual.
        abrir_gmail()

        # 4. Confirma a intenção de processar o ZIP.
        print("\nArquivo a processar:", ARQUIVO_ZIP.name)

        confirmacao = input(
            "Digite PROCESSAR para iniciar a leitura dos PDFs: "
        )

        if confirmacao != "PROCESSAR":
            print("Processamento cancelado pelo usuário.")
            return

        # 5. Lê e salva os PDFs.
        print("\nIniciando o processamento do ZIP...")

        processados, inseridos, ignorados, erros = (
            processar_zip(conexao)
        )

        # 6. Exporta os dados para Excel.
        print("\nExportando os registros para Excel...")

        exportar_excel(conexao)

        # 7. Mostra o resumo.
        print("\n" + "=" * 60)
        print("              RESUMO FINAL")
        print("=" * 60)

        print(f"PDFs com texto lido:       {processados}")
        print(f"Novos registros inseridos: {inseridos}")
        print(f"Documentos ignorados:      {ignorados}")
        print(f"Erros:                     {erros}")

        print(f"Relatório Excel: {ARQUIVO_EXCEL.name}")

        print("=" * 60)
        print("PROCESSAMENTO FINALIZADO!")

    except Exception as erro:
        print(f"\nERRO DURANTE A EXECUÇÃO: {erro}")

    finally:
        if conexao is not None and conexao.is_connected():
            conexao.close()
            print("\nConexão com o MySQL encerrada.")


# ------------------------------------------------------------
# INÍCIO DO PROGRAMA
# ------------------------------------------------------------

if __name__ == "__main__":
    main()
